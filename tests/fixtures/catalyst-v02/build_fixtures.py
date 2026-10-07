"""Build deterministic synthetic Catalyst v0.2 mixed-corpus fixtures.

This module creates real PDF, Office Open XML, raster, text, and structured
files for parser acceptance; it does not mock the parsing boundary.
本模块生成真实的 PDF、Office Open XML、图像和结构化文件，用于解析器验收，不模拟解析器。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

ROOT = Path(__file__).resolve().parent
FIXED_TIME = datetime(2026, 10, 7, tzinfo=UTC)
ZIP_TIME = (2026, 10, 7, 0, 0, 0)
MANAGEMENT_SENTINEL = "ORCHID42_OPERATOR_ONLY_SENTINEL_DO_NOT_TRAIN_20261007"
FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf"


def _normalise_zip(path: Path) -> None:
    """Rewrite a generated OOXML package with stable order and timestamps.

    统一 OOXML 包的条目顺序、时间戳和文件属性，使签名可重复。
    """
    with ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}

    path.write_bytes(_pack_zip(entries))


def _pack_zip(entries: dict[str, bytes]) -> bytes:
    """Serialize a deterministic ZIP payload from named byte entries."""
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
        for name, payload in sorted(entries.items()):
            info = ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            target.writestr(info, payload, compress_type=ZIP_DEFLATED, compresslevel=9)
    return output.getvalue()


def _fix_core_timestamps(payload: bytes) -> bytes:
    """Replace volatile OOXML core timestamps with the fixed build date."""
    fixed = b"2026-10-07T00:00:00Z"
    for tag in (b"created", b"modified"):
        pattern = rb"(<dcterms:" + tag + rb"\b[^>]*>).*?(</dcterms:" + tag + rb">)"
        payload, count = re.subn(pattern, rb"\g<1>" + fixed + rb"\g<2>", payload, count=1)
        if count != 1:
            raise RuntimeError(f"OOXML core properties lack a dcterms:{tag.decode()} value.")
    return payload


def _normalise_embedded_workbook(payload: bytes) -> bytes:
    """Normalize the generated chart workbook nested inside the PPTX."""
    with ZipFile(BytesIO(payload), "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    core_name = "docProps/core.xml"
    if core_name in entries:
        entries[core_name] = _fix_core_timestamps(entries[core_name])
    return _pack_zip(entries)


def _fixed_properties(properties: object) -> None:
    """Set common OOXML core properties to a fixed synthetic timestamp."""
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    properties.last_modified_by = "Cyrene synthetic fixture builder"
    properties.author = "Cyrene synthetic fixture builder"


def _write_images(output: Path) -> tuple[Path, Path, Path]:
    """Create clear PNG/JPEG and deliberately degraded OCR image samples.

    创建清晰 PNG/JPEG 及刻意降质的 OCR 图像样本。
    """
    from PIL import Image, ImageDraw, ImageEnhance, ImageFont

    clear = Image.new("RGB", (1600, 500), color=(250, 250, 246))
    draw = ImageDraw.Draw(clear)
    font = ImageFont.truetype(FONT_PATH, 62)
    small = ImageFont.truetype(FONT_PATH, 38)
    draw.rectangle((18, 18, 1581, 481), outline=(24, 55, 79), width=8)
    draw.text((58, 48), "IT EQUIPMENT HANDOFF", font=font, fill=(14, 32, 48))
    draw.text((60, 166), "TICKET ORCHID-42", font=font, fill=(14, 32, 48))
    draw.text((62, 300), "ASSET SP-204 / BAY 07", font=small, fill=(28, 48, 64))
    draw.text((62, 370), "STATUS RECEIVED 2026-10-07", font=small, fill=(28, 48, 64))

    clear_png = output / "handoff-ocr-clear.png"
    clear_jpeg = output / "handoff-ocr-clear.jpg"
    clear.save(clear_png, format="PNG", optimize=False, compress_level=9)
    clear.save(
        clear_jpeg,
        format="JPEG",
        quality=94,
        subsampling=0,
        optimize=False,
        progressive=False,
        dpi=(300, 300),
    )

    low_quality = clear.convert("L").resize((320, 100), Image.Resampling.BILINEAR)
    pixels = bytearray(low_quality.tobytes())
    noise = random.Random(42)
    for index in range(len(pixels)):
        pixels[index] = max(0, min(255, pixels[index] + noise.randint(-17, 17)))
    low_quality.putdata(list(pixels))
    low_quality = ImageEnhance.Contrast(low_quality).enhance(0.72)
    low_quality_jpeg = output / "handoff-ocr-low-quality.jpg"
    low_quality.save(
        low_quality_jpeg,
        format="JPEG",
        quality=12,
        subsampling=2,
        optimize=False,
        progressive=False,
    )
    return clear_png, clear_jpeg, low_quality_jpeg


def _write_native_pdf(path: Path) -> None:
    """Create a text-native PDF with selectable handoff details and a table.

    创建包含可抽取文字和步骤表格的原生文本 PDF。
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen.canvas import Canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    canvas = Canvas(str(path), pagesize=letter, pageCompression=0, invariant=1)
    canvas.setTitle("IT Equipment Handoff ORCHID-42")
    canvas.setAuthor("Cyrene synthetic fixture builder")
    canvas.setSubject("Synthetic IT equipment handoff workflow")
    canvas.setFont("Helvetica-Bold", 18)
    canvas.drawString(48, 744, "IT EQUIPMENT HANDOFF / ORCHID-42")
    canvas.setFont("STSong-Light", 13)
    canvas.drawString(48, 713, "IT 设备交接流程：归还、核对、恢复")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(48, 676, "Ticket: ORCHID-42     Asset: SP-204     Bay: 07")
    canvas.drawString(48, 653, "Synthetic training material; no personal or customer data.")

    x_positions = (48, 85, 310, 440)
    row_y = (600, 572, 544, 516)
    rows = (
        ("Step", "Action", "Evidence", "Status"),
        ("1", "Receive device", "SP-204 serial check", "complete"),
        ("2", "Verify recovery", "ORCHID-42 token", "complete"),
        ("3", "Wipe and return", "intake receipt", "pending"),
    )
    canvas.setStrokeColorRGB(0.18, 0.30, 0.40)
    for y in row_y:
        canvas.line(48, y - 7, 560, y - 7)
    for row, y in zip(rows, row_y, strict=True):
        canvas.setFont("Helvetica-Bold" if y == row_y[0] else "Helvetica", 10)
        for text, x in zip(row, x_positions, strict=True):
            canvas.drawString(x, y, text)
    canvas.showPage()
    canvas.save()


def _write_scanned_pdf(path: Path, scan_image: Path) -> None:
    """Create an image-only PDF page from the clear synthetic scan.

    从清晰合成扫描图创建仅含栅格图像的一页 PDF，不增加可抽取文字层。
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    canvas = Canvas(str(path), pagesize=letter, pageCompression=0, invariant=1)
    canvas.setTitle("Scanned IT Equipment Handoff ORCHID-42")
    canvas.setAuthor("Cyrene synthetic fixture builder")
    canvas.drawImage(
        ImageReader(str(scan_image)),
        40,
        480,
        width=532,
        height=166,
        preserveAspectRatio=True,
        mask="auto",
    )
    canvas.showPage()
    canvas.save()


def _write_docx(path: Path) -> None:
    """Write a genuine DOCX containing workflow prose and an equipment table."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    document.core_properties.title = "IT 设备交接流程 / ORCHID-42"
    document.core_properties.subject = "Synthetic device handoff business document"
    _fixed_properties(document.core_properties)
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    title = document.add_heading("IT 设备交接流程", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("流程编号：ORCHID-42；所有条目均为合成业务资料。")
    document.add_paragraph(
        "目标：核对归还设备、确认安全擦除、记录存放位置，并完成值班组之间的交接。"
    )
    document.add_heading("交接步骤", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("步骤", "设备", "交接动作", "状态")
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
    rows = (
        ("1", "SP-204", "在 Bay 07 核对资产标签", "已接收"),
        ("2", "SP-204", "使用 ORCHID-42 复核恢复流程", "已核验"),
        ("3", "SP-204", "完成擦除后移交库存", "待处理"),
    )
    for values in rows:
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = value
    paragraph = document.add_paragraph("验收备注：管理字段与审核标记不得混入训练答案。")
    paragraph.runs[0].font.size = Pt(9)
    document.save(path)
    _normalise_zip(path)


def _write_pptx(path: Path) -> dict[str, object]:
    """Write a two-slide PPTX with a notes slide, positioned text, and chart."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.core_properties.title = "IT Equipment Handoff ORCHID-42"
    presentation.core_properties.subject = "Synthetic handoff briefing"
    _fixed_properties(presentation.core_properties)

    slide_one = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_shape = slide_one.shapes.add_textbox(Inches(0.75), Inches(1.25), Inches(7.5), Inches(1.1))
    text_shape.text_frame.text = "ORCHID-42 | IT EQUIPMENT HANDOFF"
    text_shape.text_frame.paragraphs[0].font.size = Pt(26)
    text_shape.text_frame.paragraphs[0].font.bold = True
    detail = slide_one.shapes.add_textbox(Inches(0.75), Inches(2.65), Inches(6.3), Inches(1.2))
    detail.text_frame.text = "ASSET SP-204\nRETURN BAY 07\nSTATUS RECEIVED"
    detail.text_frame.paragraphs[0].font.size = Pt(18)
    accent = slide_one.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(4.1), Inches(7.5), Inches(0.08)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = __import__("pptx").dml.color.RGBColor(22, 91, 131)
    accent.line.fill.background()

    slide_two = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide_two.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(8.4), Inches(0.7))
    title.text_frame.text = "Handoff counts by receiving bay"
    title.text_frame.paragraphs[0].font.size = Pt(24)
    chart_data = CategoryChartData()
    chart_data.categories = ["Bay 07", "Bay 12", "Reserve"]
    chart_data.add_series("Devices", (3, 2, 1))
    slide_two.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.5),
        Inches(7.5),
        Inches(4.5),
        chart_data,
    )
    slide_two.notes_slide.notes_text_frame.text = (
        "Synthetic presenter note: verify ORCHID-42, confirm the asset label, then record the bay."
    )
    presentation.save(path)
    with ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    for name, payload in tuple(entries.items()):
        if name.startswith("ppt/embeddings/") and name.endswith(".xlsx"):
            entries[name] = _normalise_embedded_workbook(payload)
    path.write_bytes(_pack_zip(entries))
    _normalise_zip(path)
    return {
        "slides": 2,
        "notesSlides": 1,
        "chartParts": 1,
        "positionedText": {
            "text": "ORCHID-42 | IT EQUIPMENT HANDOFF",
            "leftEmu": int(text_shape.left),
            "topEmu": int(text_shape.top),
        },
    }


def _patch_formula_cache(path: Path) -> None:
    """Store one cached result while preserving a second uncached formula."""
    with ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    sheet_name = "xl/worksheets/sheet1.xml"
    sheet_xml = entries[sheet_name]
    cached_formula = re.compile(rb'(<c r="F2"[^>]*><f>SUM\(C2:C3\)</f><v>).*?(</v></c>)')
    sheet_xml, matches = cached_formula.subn(rb"\g<1>42\g<2>", sheet_xml, count=1)
    if matches != 1:
        raise RuntimeError("Could not set the expected cached formula value in F2.")
    entries[sheet_name] = sheet_xml
    core_name = "docProps/core.xml"
    entries[core_name] = _fix_core_timestamps(entries[core_name])
    path.write_bytes(_pack_zip(entries))


def _write_xlsx(path: Path) -> dict[str, object]:
    """Write a two-sheet workbook with a table and cached/uncached formulas."""
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    workbook.properties.title = "IT Equipment Handoff Register"
    workbook.properties.subject = "Synthetic ORCHID-42 asset inventory"
    _fixed_properties(workbook.properties)
    handoff = workbook.active
    handoff.title = "Handoff Register"
    handoff.append(("Asset", "Bay", "Units", "Status"))
    handoff.append(("SP-204", "Bay 07", 20, "received"))
    handoff.append(("SP-205", "Bay 07", 22, "verified"))
    handoff.append(("SP-206", "Bay 12", 7, "pending"))
    table = Table(displayName="HandoffTable", ref="A1:D4")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    handoff.add_table(table)
    handoff["F2"] = "=SUM(C2:C3)"
    handoff["F3"] = "=SUM(C3:C4)"
    handoff["G1"] = "Cached formula value"
    handoff["G2"] = 42

    lookup = workbook.create_sheet("Recovery Lookup")
    lookup.append(("Recovery code", "Procedure", "Owner group"))
    lookup.append(("ORCHID-42", "Verify label, wipe device, log receipt", "synthetic-itops"))
    lookup.append(("ORCHID-43", "Hold for secondary review", "synthetic-itops"))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)
    _patch_formula_cache(path)
    _normalise_zip(path)
    return {
        "sheets": ["Handoff Register", "Recovery Lookup"],
        "tables": ["HandoffTable"],
        "formulas": {
            "F2": {"formula": "SUM(C2:C3)", "cachedValue": 42},
            "F3": {"formula": "SUM(C3:C4)", "cachedValue": None},
        },
    }


def _write_text_assets(output: Path) -> dict[str, object]:
    """Write Markdown, plain text, CSV, JSONL, and deliberately invalid files."""
    (output / "handoff.md").write_text(
        "# IT 设备交接流程 / ORCHID-42\n\n"
        "全部内容为合成业务资料，不含个人或客户信息。\n\n"
        "| 步骤 | 操作 | 验收 |\n|---|---|---|\n"
        "| 1 | 核对 SP-204 与 Bay 07 | 资产标签一致 |\n"
        "| 2 | 执行 ORCHID-42 恢复流程 | 状态为 received |\n"
        "| 3 | 擦除并登记交接 | 生成合成回执 |\n",
        encoding="utf-8",
        newline="\n",
    )
    (output / "handoff.txt").write_text(
        "IT EQUIPMENT HANDOFF — ORCHID-42\n"
        "Asset: SP-204\nBay: 07\nStatus: received\n"
        "Synthetic document for mixed-corpus parser acceptance.\n",
        encoding="utf-8",
        newline="\n",
    )
    with (output / "handoff.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(("asset_id", "bay", "handoff_step", "status", "recovery_code"))
        writer.writerows(
            (
                ("SP-204", "07", "receive", "complete", "ORCHID-42"),
                ("SP-205", "07", "verify", "complete", "ORCHID-42"),
                ("SP-206", "12", "wipe", "pending", "ORCHID-43"),
            )
        )

    records: list[dict[str, object]] = []
    for family_index in range(1, 13):
        family = f"family-{family_index:02d}"
        for conversation_index in range(1, 3):
            conversation = f"ORCHID-{family_index:02d}-CONV-{conversation_index:02d}"
            records.append(
                {
                    "sampleId": f"{conversation}-SAMPLE",
                    "sourceFamily": family,
                    "conversationId": conversation,
                    "messages": [
                        {
                            "role": "user",
                            "content": (
                                f"设备 SP-{200 + family_index} 在 Bay 07 如何按 ORCHID-42 交接？"
                            ),
                        },
                        {
                            "role": "assistant",
                            "content": ("核对资产标签，记录接收状态，完成擦除后登记去向。"),
                        },
                        {
                            "role": "user",
                            "content": f"{conversation} 的最终状态是什么？",
                        },
                        {
                            "role": "assistant",
                            "content": "合成状态为 received；未完成擦除时标为 pending。",
                        },
                    ],
                    "_reviewNote": MANAGEMENT_SENTINEL,
                    "_acl": ["org:synthetic-itops"],
                    "_operatorAnnotation": "metadata only; do not train as answer text",
                    "_ingestBatch": "fixture-batch-20261007",
                }
            )
    (output / "source-conversations.jsonl").write_text(
        "".join(
            json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n"
            for record in records
        ),
        encoding="utf-8",
        newline="\n",
    )
    (output / "malformed.pdf").write_bytes(
        b"%PDF-1.7\n1 0 obj << /Type /Catalog /Pages 2 0 R >>\n"
        b"this fixture intentionally has no valid xref or EOF marker\n"
    )
    (output / "unsupported.rtf").write_text(
        r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Arial;}}\f0\fs24 Synthetic unsupported RTF ORCHID-42\par}",
        encoding="ascii",
        newline="\n",
    )
    return {
        "records": len(records),
        "sourceFamilies": 12,
        "conversations": 24,
        "managementFields": ["_acl", "_ingestBatch", "_operatorAnnotation", "_reviewNote"],
        "managementSentinel": MANAGEMENT_SENTINEL,
    }


def _file_record(
    output: Path,
    name: str,
    media_type: str,
    expected_structure: dict[str, object],
    warning_expectation: list[str] | str,
) -> dict[str, object]:
    """Return a fixed digest and the expected structure for one asset."""
    payload = (output / name).read_bytes()
    return {
        "sha256": hashlib.sha256(payload).hexdigest(),
        "sizeBytes": len(payload),
        "mediaType": media_type,
        "expectedStructure": expected_structure,
        "warningExpectation": warning_expectation,
    }


def build(output: Path = ROOT) -> dict[str, object]:
    """Build all deterministic fixtures in ``output`` and return the manifest.

    生成固定混合语料和签名清单；可传临时目录验证可重复生成。
    """
    output.mkdir(parents=True, exist_ok=True)
    clear_png, clear_jpeg, low_jpeg = _write_images(output)
    native_pdf = output / "handoff-native.pdf"
    scanned_pdf = output / "handoff-scanned.pdf"
    _write_native_pdf(native_pdf)
    _write_scanned_pdf(scanned_pdf, clear_png)
    _write_docx(output / "handoff.docx")
    pptx_structure = _write_pptx(output / "handoff.pptx")
    xlsx_structure = _write_xlsx(output / "handoff.xlsx")
    source_structure = _write_text_assets(output)

    files: dict[str, dict[str, object]] = {
        "handoff-native.pdf": _file_record(
            output,
            "handoff-native.pdf",
            "application/pdf",
            {"pages": 1, "nativeText": True, "sentinels": ["ORCHID-42", "SP-204"]},
            [],
        ),
        "handoff-scanned.pdf": _file_record(
            output,
            "handoff-scanned.pdf",
            "application/pdf",
            {
                "pages": 1,
                "nativeText": False,
                "imageObjects": 1,
                "sourceImage": "handoff-ocr-clear.png",
            },
            ["OCR is required because the page has no native text layer."],
        ),
        "handoff.docx": _file_record(
            output,
            "handoff.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            {"paragraphsAtLeast": 5, "tables": 1, "sentinels": ["ORCHID-42", "SP-204"]},
            [],
        ),
        "handoff.pptx": _file_record(
            output,
            "handoff.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            pptx_structure,
            (
                "PPTX is supported, but chart object structure and data are not fully "
                "extracted, so an explicit warning is expected."
            ),
        ),
        "handoff.xlsx": _file_record(
            output,
            "handoff.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_structure,
            "F3 deliberately has no cached formula value; an explicit cache warning is expected.",
        ),
        "handoff-ocr-clear.png": _file_record(
            output,
            clear_png.name,
            "image/png",
            {
                "width": 1600,
                "height": 500,
                "quality": "clear",
                "ocrSentinels": ["ORCHID-42", "SP-204"],
            },
            [],
        ),
        "handoff-ocr-clear.jpg": _file_record(
            output,
            clear_jpeg.name,
            "image/jpeg",
            {
                "width": 1600,
                "height": 500,
                "quality": "clear",
                "ocrSentinels": ["ORCHID-42", "SP-204"],
            },
            [],
        ),
        "handoff-ocr-low-quality.jpg": _file_record(
            output,
            low_jpeg.name,
            "image/jpeg",
            {
                "width": 320,
                "height": 100,
                "quality": "low",
                "ocrSentinels": ["ORCHID-42", "SP-204"],
            },
            ["Low resolution and compression artifacts may reduce OCR confidence."],
        ),
        "handoff.md": _file_record(
            output,
            "handoff.md",
            "text/markdown",
            {"heading": "IT 设备交接流程 / ORCHID-42", "tableRows": 3},
            [],
        ),
        "handoff.txt": _file_record(
            output,
            "handoff.txt",
            "text/plain",
            {"sentinels": ["ORCHID-42", "SP-204", "Bay: 07"]},
            [],
        ),
        "handoff.csv": _file_record(
            output,
            "handoff.csv",
            "text/csv",
            {"rows": 3, "columns": 5, "sentinels": ["ORCHID-42", "SP-204"]},
            [],
        ),
        "source-conversations.jsonl": _file_record(
            output,
            "source-conversations.jsonl",
            "application/x-ndjson",
            source_structure,
            ["Management fields must remain metadata and must not become answer text."],
        ),
        "malformed.pdf": _file_record(
            output,
            "malformed.pdf",
            "application/pdf",
            {"validPdf": False, "reason": "missing xref and EOF marker"},
            ["Invalid PDF structure should fail independently of unsupported input types."],
        ),
        "unsupported.rtf": _file_record(
            output,
            "unsupported.rtf",
            "application/rtf",
            {"validRtfText": True, "sentinel": "ORCHID-42"},
            ["RTF is intentionally outside the supported v0.2 formats."],
        ),
    }
    manifest: dict[str, object] = {
        "schemaVersion": 1,
        "purpose": "Synthetic fixed mixed corpus for Catalyst v0.2 parser acceptance.",
        "topic": "IT equipment handoff workflow / ORCHID-42",
        "synthetic": True,
        "personalOrCustomerData": False,
        "managementSentinel": MANAGEMENT_SENTINEL,
        "files": files,
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    return manifest


def main() -> None:
    """Run the fixture builder for its checked-in directory or a temp directory."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT, help="target directory")
    args = parser.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
